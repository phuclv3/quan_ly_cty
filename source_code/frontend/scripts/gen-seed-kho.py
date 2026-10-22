"""
Sinh src/app/core/seed-kho.ts từ tai_lieu/note_update/KhoOxyMecie_v3.xlsx.

Chạy (từ thư mục source_code/frontend):
    python -X utf8 scripts/gen-seed-kho.py ../../tai_lieu/note_update/KhoOxyMecie_v3.xlsx

Sheet dùng:
  - "DM. Mã hàng": danh mục mã hàng CỐ ĐỊNH (mã, nhóm, loại khí / vật liệu, thể tích, ĐVT, tên tham chiếu).
    Cột "ANH SỬA" có giá trị → dùng làm mã thay cho cột "MÃ HÀNG".
  - "1. NhapKho": mỗi dòng hoá đơn = một LÔ (mã lô, NCC, số hoá đơn, quy cách / nồng độ, SL, đơn giá chưa VAT, VAT).
  - "6. KiemKe": kiểm kê thực tế. Cột "MÃ HÀNG (anh gán)" trống → gán tự động theo tên + thể tích,
    dòng không gán được thì bỏ qua (in ra màn hình để kiểm tra).
"""
import json
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / 'src' / 'app' / 'core' / 'seed-kho.ts'

# Nhóm trong file Excel → nhóm hàng của app (NHOM_HANG).
NHOM = {'Thiết bị': 'Thiết bị tái sử dụng'}


def s(v):
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def rows(ws, start):
    for r in ws.iter_rows(min_row=start, values_only=True):
        if any(v is not None and str(v).strip() for v in r):
            yield r


def ma_hang(ws):
    out = []
    for r in rows(ws, 5):
        ma = s(r[1]) or s(r[0])
        if not ma:
            continue
        bd = [x.strip() for x in s(r[6]).split('|') if x.strip()]
        o = {'ma': ma, 'nhom': NHOM.get(s(r[2]), s(r[2])), 'dvt': s(r[5]) or 'Bình'}
        if s(r[3]):
            o['loaiKhi'] = s(r[3])
        if s(r[4]):
            o['theTich'] = s(r[4])
        if bd:
            o['bietDanh'] = bd
        out.append(o)
    return out


def nhap_kho(ws):
    out = []
    for r in rows(ws, 5):
        if not s(r[2]):
            continue
        out.append({
            'ngay': s(r[0]),
            'soHoaDon': s(r[1]),
            'lo': s(r[2]),
            'ncc': s(r[3]),
            'ma': s(r[5]),
            'quyCach': s(r[6]),
            'soLuong': int(r[7] or 0),
            'donGia': int(r[8] or 0),
            'vat': round(float(r[10] or 0) * 100),
        })
    return out


def gan_ma(nhom, ten, the_tich, cac_ma):
    """Kiểm kê: "Khí chuẩn" + "SO2/NO/CO" + "8L" → KC-SO2NOCO-8L; vỏ nhôm/thép/khác theo thể tích."""
    vol = the_tich.upper().replace(' ', '')
    if nhom == 'Vỏ bình':
        t = ten.lower()
        loai = 'NHOM' if 'nhôm' in t else 'THEP' if 'thép' in t else 'KHAC'
        for c in (f'VO-{loai}-{vol}', f'VO-{loai}', f'VO-KHAC-{vol}'):
            if c in cac_ma:
                return c
        return ''
    if ten == 'Toàn Tiếng TQ':
        khi = 'TO'
    else:
        khi = re.sub(r'[^A-Z0-9]', '', ten.upper())
    for c in (f'KC-{khi}-{vol}', f'KC-{khi}'):
        if c in cac_ma:
            return c
    return ''


def kiem_ke(ws, cac_ma):
    ngay = ''
    m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', s(ws.cell(1, 1).value))
    if m:
        ngay = f'{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}'
    dong = []
    for r in rows(ws, 5):
        nhom, ten, xx, tt, nd, sl, gan = (s(x) for x in r[:7])
        if not nhom or not sl:
            continue
        ma = gan or gan_ma(nhom, ten, tt, cac_ma)
        if not ma:
            print(f'  ! Kiểm kê: không gán được mã cho "{nhom} / {ten} / {tt}" × {sl}')
            continue
        o = {'ma': ma, 'soLuong': int(float(sl)), 'tenHang': ten}
        if xx:
            o['xuatXu'] = xx
        if nd:
            o['nongDo'] = nd
        dong.append(o)
    return {'ngay': ngay, 'dong': dong}


def anh_xa(ws, cac_ma):
    """Sheet "5. AnhXaTenBan": tên hàng trong HĐ bán → mã hàng / mã lô (gán tay)."""
    out = []
    for r in rows(ws, 5):
        ten = s(r[0])
        if not ten:
            continue
        o = {'ten': ten, 'tongSL': float(r[3] or 0), 'soLan': int(r[4] or 0)}
        for k, i in (('theTich', 1), ('nongDo', 2), ('lo', 6), ('ghiChu', 7)):
            if s(r[i]):
                o[k] = s(r[i])
        if s(r[5]):
            if s(r[5]) in cac_ma:
                o['ma'] = s(r[5])
            else:
                print(f'  ! Ánh xạ: mã "{s(r[5])}" của "{ten}" không có trong danh mục')
        out.append(o)
    return out


def main():
    if len(sys.argv) < 2:
        sys.exit('Thiếu đường dẫn file Excel')
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    mh = ma_hang(wb['DM. Mã hàng'])
    cac_ma = {x['ma'] for x in mh}
    nk = nhap_kho(wb['1. NhapKho'])
    lech = sorted({x['ma'] for x in nk} - cac_ma)
    if lech:
        print('  ! Mã hàng ở NhapKho không có trong danh mục:', lech)
    kk = kiem_ke(wb['6. KiemKe'], cac_ma)
    ax = anh_xa(wb['5. AnhXaTenBan'], cac_ma)
    j = lambda o: json.dumps(o, ensure_ascii=False, indent=1)
    OUT.write_text(
        '/**\n'
        ' * Dữ liệu kho — trích từ tai_lieu/note_update/KhoOxyMecie_v3.xlsx (sheet "DM. Mã hàng", "1. NhapKho", "5. AnhXaTenBan", "6. KiemKe").\n'
        ' * Tệp sinh tự động bằng scripts/gen-seed-kho.py — không sửa tay.\n'
        ' */\n'
        'export interface RawMaHang {\n  ma: string;\n  nhom: string;\n  dvt: string;\n  loaiKhi?: string;\n  theTich?: string;\n  bietDanh?: string[];\n}\n\n'
        'export interface RawLoNhap {\n  ngay: string;\n  soHoaDon: string;\n  lo: string;\n  ncc: string;\n  ma: string;\n  quyCach: string;\n  soLuong: number;\n  donGia: number;\n  vat: number;\n}\n\n'
        'export interface RawKiemKe {\n  ngay: string;\n  dong: { ma: string; soLuong: number; tenHang: string; xuatXu?: string; nongDo?: string }[];\n}\n\n'
        'export interface RawAnhXa {\n  ten: string;\n  theTich?: string;\n  nongDo?: string;\n  tongSL: number;\n  soLan: number;\n  ma?: string;\n  lo?: string;\n  ghiChu?: string;\n}\n\n'
        f'export const RAW_MA_HANG: RawMaHang[] = {j(mh)};\n\n'
        f'export const RAW_LO_NHAP: RawLoNhap[] = {j(nk)};\n\n'
        f'export const RAW_KIEM_KE: RawKiemKe = {j(kk)};\n\n'
        f'export const RAW_ANH_XA: RawAnhXa[] = {j(ax)};\n',
        encoding='utf-8',
    )
    print(f'Đã ghi {OUT} — {len(mh)} mã hàng, {len(nk)} lô, {len(kk["dong"])} dòng kiểm kê, {len(ax)} tên ánh xạ')


if __name__ == '__main__':
    main()
